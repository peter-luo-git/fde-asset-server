"""附件图片理解：规则筛选 → 大模型按文字挑 → 只理解挑中的图 → 说明插回文字版（2026-10-28）。"""

from __future__ import annotations

import hashlib
import io
import json
from types import SimpleNamespace

from fde_asset.modules.asset import attachment_vision as vision


def _png(color: tuple[int, int, int], size: tuple[int, int] = (400, 300)) -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", size, color).save(out, "PNG")
    return out.getvalue()


def _deck() -> bytes:
    """5 页：每页右下角同一个 logo；第 2 页一张架构图，第 4 页一张图表；第 3 页一个小图标。"""
    from pptx import Presentation
    from pptx.util import Emu

    presentation = Presentation()
    width, height = presentation.slide_width, presentation.slide_height
    logo = _png((200, 0, 0), (300, 300))
    for number in range(1, 6):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"第{number}页标题"
        slide.shapes.add_picture(
            io.BytesIO(logo),
            Emu(int(width * 0.9)),
            Emu(int(height * 0.93)),
            Emu(int(width * 0.05)),
            Emu(int(height * 0.05)),
        )
        if number == 2:
            slide.shapes.add_picture(
                io.BytesIO(_png((0, 0, 200))),
                Emu(int(width * 0.1)),
                Emu(int(height * 0.3)),
                Emu(int(width * 0.6)),
                Emu(int(height * 0.5)),
            )
        if number == 3:
            slide.shapes.add_picture(
                io.BytesIO(_png((0, 200, 0), (60, 60))),
                Emu(int(width * 0.4)),
                Emu(int(height * 0.4)),
                Emu(int(width * 0.05)),
                Emu(int(height * 0.05)),
            )
        if number == 4:
            slide.shapes.add_picture(
                io.BytesIO(_png((0, 120, 120))),
                Emu(int(width * 0.2)),
                Emu(int(height * 0.3)),
                Emu(int(width * 0.5)),
                Emu(int(height * 0.5)),
            )
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


class FakeLlm:
    usable = True

    def __init__(self, skip: set[str] | None = None) -> None:
        self.calls: list[tuple[str, object]] = []
        self.skip = skip or set()

    def complete_json(self, system, user):
        self.calls.append((system, user))
        if isinstance(user, str):  # 按文字挑
            pages = json.loads(user)["pages"]
            ids = [image["id"] for page in pages for image in page["images"]]
            return {
                "selected": [{"id": i, "reason": "有内容"} for i in ids if i not in self.skip],
                "skipped": [{"id": i, "reason": "封面配图"} for i in ids if i in self.skip],
            }
        text = user[0]["text"]
        return {"decorative": False, "description": f"图里是{text.split('》')[1][:5]}的架构"}


def test_rules_drop_repeated_logo_and_tiny_icon() -> None:
    units = vision.outline("deck.pptx", _deck())
    skipped = vision.rule_filter(units)
    assert len(units) == 5
    logos = [p.id for u in units for p in u.pictures if p.id.endswith("-i1")]
    assert all("重复" in skipped[i] or "页眉页脚" in skipped[i] for i in logos)
    assert any("太小" in reason for reason in skipped.values())
    kept = [p.id for u in units for p in u.pictures if p.id not in skipped]
    assert kept == ["p2-i2", "p4-i2"]


def test_understand_inserts_descriptions_in_order_and_reports_skips() -> None:
    llm = FakeLlm(skip={"p4-i2"})
    progress: list[tuple[int, int]] = []
    text, report = vision.understand(
        "deck.pptx", _deck(), llm, progress=lambda d, t: progress.append((d, t))
    )
    # 只看了第 2 页那张：logo、图标被规则筛掉，第 4 页那张被大模型按文字跳过
    looked = [user for _, user in llm.calls if not isinstance(user, str)]
    assert len(looked) == 1
    assert "【第 2 页图片说明】图里是第 2 页的架构" in text
    assert text.index("第2页标题") < text.index("【第 2 页图片说明】") < text.index("第3页标题")
    assert {entry["id"] for entry in report["described"]} == {"p2-i2"}
    assert {"id": "p4-i2", "reason": "封面配图"} in report["skipped"]
    assert progress[-1] == (1, 1)


def test_include_forces_a_skipped_picture() -> None:
    llm = FakeLlm(skip={"p4-i2"})
    text, report = vision.understand("deck.pptx", _deck(), llm, include={"p4-i2"})
    assert "【第 4 页图片说明】" in text


def test_single_image_attachment_gets_a_text_version() -> None:
    text, report = vision.understand("架构.png", _png((10, 10, 10)), FakeLlm())
    assert "图片说明" in text and report["calls"] == 1


def test_queue_runs_once_per_file_and_feeds_the_agent_text(indexed) -> None:
    from fde_asset.modules.asset.snapshot import with_attachment_texts

    data = _deck()
    sha = indexed.file_store.put(data)
    assert sha == hashlib.sha256(data).hexdigest()
    vision.enqueue(indexed.engine, sha, "deck.pptx")
    vision.enqueue(indexed.engine, sha, "deck.pptx")  # 同一份只排一次
    assert vision.run_one(indexed.engine, indexed.file_store, lambda: FakeLlm())["status"] == "done"
    assert vision.run_one(indexed.engine, indexed.file_store, lambda: FakeLlm()) == {}
    saved = vision.get(indexed.engine, sha)
    assert saved["status"] == "done" and saved["done"] == saved["total"] == 2

    row = SimpleNamespace(
        name="x",
        attachments_json=json.dumps(
            [
                {
                    "path": "attachments/deck.pptx",
                    "name": "deck.pptx",
                    "sha256": sha,
                    "text_extraction": "ok",
                }
            ]
        ),
    )
    assert vision.texts_for(indexed.engine, [row]) == {sha: saved["text"]}
    readme, texts, _ = with_attachment_texts(
        row, b"# x\n", indexed.file_store, enriched=vision.texts_for(indexed.engine, [row])
    )
    assert "图片说明" in texts["attachments-text/deck.pptx.txt"].decode("utf-8")


def test_without_model_nothing_is_claimed(indexed) -> None:
    sha = indexed.file_store.put(_png((1, 2, 3)))
    vision.enqueue(indexed.engine, sha, "a.png")
    assert vision.run_one(indexed.engine, indexed.file_store, lambda: None) == {
        "status": "waiting_model"
    }
    assert vision.get(indexed.engine, sha)["status"] == "pending"


def test_texts_for_queues_attachments_landed_before(indexed) -> None:
    """以前入库、还没做过图片理解的附件，拼资产包时顺手排上。"""
    sha = indexed.file_store.put(_png((4, 5, 6)))
    row = SimpleNamespace(
        name="old",
        attachments_json=json.dumps(
            [{"path": "attachments/图.png", "name": "图.png", "sha256": sha}]
        ),
    )
    assert vision.texts_for(indexed.engine, [row]) == {}
    assert vision.get(indexed.engine, sha)["status"] == "pending"


def test_draft_attachment_text_version_api(client) -> None:
    """草稿上传 PPT 后排上图片理解；页面能查进度、重新理解、点「这页也要理解」。"""
    import base64

    client.headers.update({"X-FDE-User": "chen"})
    draft = client.post(
        "/api/v1/harvest-candidates",
        json={"kind": "Solution", "title": "带图的方案", "scope": "company"},
    ).json()
    candidate_id = draft["candidate_id"]
    added = client.post(
        f"/api/v1/harvest-candidates/{candidate_id}/attachments",
        json={"filename": "deck.pptx", "content_base64": base64.b64encode(_deck()).decode()},
    )
    assert added.status_code == 200, added.text
    path = f"/api/v1/harvest-candidates/{candidate_id}/attachments/deck.pptx/text-version"
    version = client.get(path).json()
    assert version["status"] == "pending"
    retried = client.post(path + "/retry", json={"target": "p4-i2"}).json()
    assert retried["status"] == "pending"
    # 规范不收附件、技能附件不做图片理解；文字文件不适用
    client.post(
        f"/api/v1/harvest-candidates/{candidate_id}/attachments",
        json={"filename": "说明.md", "content_base64": base64.b64encode("文字".encode()).decode()},
    )
    other = client.get(
        f"/api/v1/harvest-candidates/{candidate_id}/attachments/说明.md/text-version"
    )
    assert other.json() == {"status": "not_applicable"}


def test_search_index_uses_text_with_picture_descriptions() -> None:
    """资产中心检索用带图片说明的文字版：截图里的内容也能搜到；图片附件也有了文字（2026-10-28）。"""
    from fde_asset.modules.asset.indexer import _attachments

    deck, picture = _deck(), _png((7, 8, 9))
    blobs = {
        hashlib.sha256(deck).hexdigest(): deck,
        hashlib.sha256(picture).hexdigest(): picture,
    }
    manifest = SimpleNamespace(
        spec=SimpleNamespace(
            attachments=[
                SimpleNamespace(name="deck.pptx", sha256=sha, bytes=len(data), summary="")
                if data is deck
                else SimpleNamespace(name="架构.png", sha256=sha, bytes=len(data), summary="")
                for sha, data in blobs.items()
            ]
        )
    )
    store = SimpleNamespace(get=blobs.get)
    enriched = {sha: f"【第 1 页图片说明】华电燃机运维大屏 {i}" for i, sha in enumerate(blobs)}
    listed, chunks = _attachments(manifest, store, 200_000, enriched)
    assert all(
        entry["text_extraction"] == "ok" and entry["detail"] == "含图片说明" for entry in listed
    )
    assert sum("华电燃机运维大屏" in chunk for chunk in chunks) == 2
    # 没做图片理解的照旧：PPT 提取普通文字，图片提取不了
    plain, plain_chunks = _attachments(manifest, store, 200_000)
    statuses = {entry["name"]: entry["text_extraction"] for entry in plain}
    assert statuses == {"deck.pptx": "ok", "架构.png": "unsupported"}
