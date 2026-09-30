"""附件文本提取：PDF / Word / Excel / PPT / Markdown，纯 Python 库，不引入服务。"""

from __future__ import annotations

import io
from dataclasses import dataclass

TEXT_SUFFIXES = {".md", ".txt", ".csv", ".json", ".yaml", ".yml"}


@dataclass(frozen=True)
class Extraction:
    text: str
    status: str  # ok | unsupported | failed
    detail: str = ""


def extract_text(filename: str, data: bytes, limit: int = 200_000) -> Extraction:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if suffix in TEXT_SUFFIXES:
            return Extraction(data.decode("utf-8", errors="replace")[:limit], "ok")
        if suffix == ".pdf":
            return _pdf(data, limit)
        if suffix == ".docx":
            return _docx(data, limit)
        if suffix == ".xlsx":
            return _xlsx(data, limit)
        if suffix == ".pptx":
            return _pptx(data, limit)
    except Exception as exc:  # noqa: BLE001 - 提取失败不应影响资产有效性
        return Extraction("", "failed", f"{type(exc).__name__}: {exc}")
    return Extraction("", "unsupported", f"暂不支持的格式：{suffix or '无扩展名'}")


def _pdf(data: bytes, limit: int) -> Extraction:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    chunks = [(page.extract_text() or "") for page in reader.pages]
    text = "\n".join(chunks).strip()
    if not text:
        return Extraction("", "unsupported", "PDF 没有文本层（可能是扫描件），v0.1 不做 OCR")
    return Extraction(text[:limit], "ok")


def _docx(data: bytes, limit: int) -> Extraction:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return Extraction("\n".join(parts)[:limit], "ok")


def _xlsx(data: bytes, limit: int, max_rows: int = 500) -> Extraction:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"[{sheet.title}]")
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= max_rows:
                parts.append("...（已截断）")
                break
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return Extraction("\n".join(parts)[:limit], "ok")


def _pptx(data: bytes, limit: int) -> Extraction:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts.append(f"[第 {index} 页]")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
    return Extraction("\n".join(parts)[:limit], "ok")
