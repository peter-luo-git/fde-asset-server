"""资产示范库的内容：FCP 平台预置的优质 SOP 和技能（2026-10-26）。

每份资产是一个目录：{相对路径: 文本}。asset.yaml 由 showcase.py 按下面的元数据生成。
SOP 的步骤用「技能：」点名配套的技能——任务按流程执行时，这一步点名的技能会自动带给 Agent。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

SHOWCASE_TAG = "示范资产"
OWNER = "department:platform"

# ---------- SOP ----------

DOC_SOP = """# 技术文档编写流程（示范）

## 结论与目的
写一份别人能独立看懂、能照着做的技术文档：先定读者和范围，再写初稿，最后评审定稿。

## 适用范围
设计文档、方案说明、操作手册、接口说明等需要交给别人阅读的技术文档。

## 角色
文档作者负责编写，文档负责人确认范围和评审结论。

## 前置条件
项目仓库可以提交文件。

## 步骤
1. 明确读者与范围
   写清这份文档给谁看、解决什么问题、不包括什么，列出大纲。
   - 检查点[交付物]：`docs/文档大纲.md` 已提交，含「读者」「范围」「大纲」三节
   - 检查点[确认:文档负责人]：负责人确认读者和范围
   - 产出：文档大纲
   - 技能：fcp-tech-doc-writer
2. 写初稿
   按大纲写正文，结论先行，每节回答一个问题；术语第一次出现时解释。
   - 检查点[交付物]：`docs/技术文档.md` 已提交，含「背景」「方案」「风险」三节
   - 检查点[命令]：正文没有残留的待补充标记 `python3 -c "import pathlib,sys; t=pathlib.Path('docs/技术文档.md').read_text(encoding='utf-8'); sys.exit(1 if ('TODO' in t or '待补充' in t) else 0)"`
   - 产出：文档初稿
   - 技能：fcp-tech-doc-writer
3. 评审与定稿
   请读者按评审清单通读，处理全部意见后定稿。
   - 检查点[判断]：不看作者解释，读者能独立看懂并照着做
   - 检查点[确认:文档负责人]：评审意见都已处理
   - 产出：定稿文档

## 异常处理
范围不清时回到第 1 步，不要边写边改范围。

## 变更记录
- 2026-10-26 FCP 示范资产初版
"""

CODE_SOP = """# 代码改动交付流程（示范）

## 结论与目的
一次代码改动从需求到合并的最小闭环：先说清楚改什么，再实现并自测，最后按清单自查、走评审。

## 适用范围
功能开发、缺陷修复、重构等需要合并进主干的代码改动。

## 角色
开发者负责实现和自测，评审人负责代码评审。

## 前置条件
项目有代码仓库，能本地运行测试。

## 步骤
1. 理解需求
   写清楚这次改动的目标、改动范围和怎么验证，范围外的事不做。
   - 检查点[交付物]：`docs/改动说明.md` 已提交，含「目标」「改动范围」「验证方式」三节
   - 检查点[确认:负责人]：负责人确认改动范围
   - 产出：改动说明
2. 实现与自测
   小步提交，每一步都能运行；新增逻辑要有单元测试。
   - 检查点[命令]：单元测试全部通过 `python3 -m pytest -q`
   - 检查点[判断]：按评审清单自查，没有遗留的调试代码和待办标记
   - 产出：代码和测试
   - 技能：fcp-code-review-checklist
3. 评审与合并
   提交合并请求，写清改了什么、怎么验证；处理评审意见后合并。
   - 检查点[确认:评审人]：合并请求已通过评审
   - 产出：已合并的改动
   - 技能：fcp-code-review-checklist

## 异常处理
测试不过不提交评审；评审发现范围外的问题另开任务。

## 变更记录
- 2026-10-26 FCP 示范资产初版
"""

SLIDES_SOP = """# 汇报材料（PPT）制作流程（示范）

## 结论与目的
做一份结论先行、每页一个观点的汇报材料：先定受众和核心结论，再写页面，最后生成可编辑的 PPT 并走查定稿。

## 适用范围
项目周报、方案汇报、复盘分享等需要对外讲解的材料。

## 角色
汇报人负责内容和定稿。

## 前置条件
知道汇报对象和时长。

## 步骤
1. 明确受众与核心结论
   一句话写清：讲给谁、希望他们记住什么、听完要做什么决定。
   - 检查点[交付物]：`slides/提纲.md` 已提交，含「受众」「核心结论」「页面提纲」三节
   - 检查点[确认:汇报人]：汇报人确认核心结论
   - 产出：汇报提纲
   - 技能：fcp-markdown-slides
2. 写页面草稿
   用 Markdown 写页面（每页用 `---` 分隔），每页标题就是这一页的结论。
   - 检查点[交付物]：`slides/汇报.md` 已提交
   - 检查点[命令]：页数不超过 15 页 `python3 -c "import pathlib,sys; t=pathlib.Path('slides/汇报.md').read_text(encoding='utf-8'); sys.exit(0 if t.count('\\n---') <= 15 else 1)"`
   - 产出：页面草稿
   - 技能：fcp-markdown-slides
3. 生成 PPT 并走查定稿
   把页面稿转成可编辑的 PPT，从头到尾讲一遍，删掉和核心结论无关的页。
   - 检查点[交付物]：`slides/汇报.pptx` 已生成
   - 检查点[命令]：PPT 能打开且页数在 2～16 页之间 `python3 -c "import sys; from pptx import Presentation; n=len(Presentation('slides/汇报.pptx').slides); sys.exit(0 if 2 <= n <= 16 else 1)"`
   - 检查点[判断]：每页一个观点、标题即结论、数据有出处
   - 检查点[确认:汇报人]：汇报人确认定稿
   - 产出：定稿 PPT
   - 技能：fcp-pptx-builder

## 异常处理
时长不够时先砍细节页，不砍结论页。

## 变更记录
- 2026-10-26 FCP 示范资产初版
"""

# ---------- 技能 ----------

DOC_SKILL = """---
name: fcp-tech-doc-writer
description: 在写设计文档、方案说明、操作手册等技术文档时使用：按读者和范围定结构，结论先行，给出模板和自查清单
---

# 技术文档写作

## 适用场景
要写一份交给别人看的技术文档，或者要把零散的说明整理成文档。

## 步骤
1. 先问清三件事：读者是谁、读完要能做什么、不包括什么。写不清就先问用户，不要猜。
2. 按 `references/文档模板.md` 搭结构，删掉用不上的节，不要硬凑。
3. 每节开头一句话说结论，再展开；一段只讲一件事。
4. 术语第一次出现时解释；缩写写全称。
5. 有步骤的写成编号列表，每步一个动作，能照着做。
6. 写完按下面的清单自查，再交给用户。

## 自查清单
- 标题能看出这份文档讲什么
- 第一屏就能看到结论和适用范围
- 没有「TODO」「待补充」之类的残留
- 每个链接、命令、路径都真实存在
- 风险和不适用的情况写清楚了

## 注意事项
- 不要堆形容词（「非常」「完善」「强大」），用事实和数字说话。
- 不确定的地方明确写「待确认：……」，并告诉用户。
"""

DOC_TEMPLATE = """# <文档标题>

> 读者：<给谁看> · 状态：草稿 / 评审中 / 定稿 · 更新：<日期>

## 结论
<一两句话：这份文档的核心结论或建议>

## 背景
<为什么要做这件事，现状和问题>

## 范围
- 包括：<…>
- 不包括：<…>

## 方案
<怎么做，关键设计和取舍>

## 风险
| 风险 | 影响 | 应对 |
|---|---|---|
| <…> | <…> | <…> |

## 附录
<术语表、参考链接>
"""

REVIEW_SKILL = """---
name: fcp-code-review-checklist
description: 在提交代码前自查或评审别人的代码改动时使用：按清单逐项检查正确性、测试、可读性和安全，并用脚本找出遗留的调试代码和待办标记
---

# 代码评审清单

## 适用场景
提交合并请求之前自查，或评审别人的改动。

## 步骤
1. 先读改动说明，确认这次改动要解决什么，范围外的问题另记。
2. 运行 `python3 scripts/find_leftovers.py <要检查的目录>`，找出遗留的 TODO、FIXME、调试打印和断点。
3. 按下面的清单逐项看，每项给出「通过 / 有问题（附位置）」。
4. 汇总：先列必须改的，再列建议改的，每条写清文件和行号。

## 清单
- 正确性：边界条件、空值、并发、错误处理有没有覆盖
- 测试：新增逻辑有测试；测试名能看出测的是什么
- 可读性：命名直白；函数只做一件事；没有大段注释掉的代码
- 安全：没有硬编码密钥、口令；外部输入有校验
- 兼容：接口和数据格式的改动有没有影响调用方

## 注意事项
- 只评这次改动，不顺手重写不相关的代码。
- 拿不准的写成问题问作者，不要直接下结论。
"""

REVIEW_SCRIPT = '''#!/usr/bin/env python3
"""找出遗留的待办标记、调试打印和断点。用法：python3 find_leftovers.py [目录]"""

import pathlib
import re
import sys

PATTERNS = [
    ("待办标记", re.compile(r"\\b(TODO|FIXME|XXX)\\b")),
    ("调试打印", re.compile(r"^\\s*(print\\(|console\\.log\\()")),
    ("断点", re.compile(r"\\b(breakpoint\\(\\)|pdb\\.set_trace\\(\\)|debugger;)")),
]
SUFFIXES = {".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".go", ".rb", ".sh"}
SKIP = {".git", "node_modules", ".venv", "dist", "build", "__pycache__"}


def main() -> int:
    root = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else ".")
    found = 0
    for path in sorted(root.rglob("*")):
        if path.suffix not in SUFFIXES or any(part in SKIP for part in path.parts) or not path.is_file():
            continue
        for number, line in enumerate(path.read_text(encoding="utf-8", errors="ignore").splitlines(), 1):
            for label, pattern in PATTERNS:
                if pattern.search(line):
                    print(f"{path}:{number}: {label}: {line.strip()[:120]}")
                    found += 1
    print(f"共 {found} 处" if found else "没有发现遗留的待办、调试打印和断点")
    return 0


if __name__ == "__main__":
    sys.exit(main())
'''

SLIDES_SKILL = """---
name: fcp-markdown-slides
description: 在需要做汇报 PPT、分享材料或演示页面时使用：先写提纲再用 Markdown 写页面，每页标题即结论，可以用脚本把提纲转成页面骨架
---

# 用 Markdown 写汇报材料

## 适用场景
项目汇报、方案讲解、复盘分享，需要一份结构清楚的演示材料。

## 步骤
1. 先问清：受众是谁、讲多久、希望他们记住哪一句话。
2. 写提纲（`slides/提纲.md`）：一级标题是核心结论，二级标题是每一页的结论。
3. 运行 `python3 scripts/outline_to_slides.py slides/提纲.md slides/汇报.md`，把提纲转成页面骨架（每页用 `---` 分隔）。
4. 按 `templates/页面模板.md` 填每一页：标题是结论，正文最多 3 条要点，数据写出处。
5. 页数控制在 15 页以内；第一页给结论，最后一页给需要对方做的决定。
6. 要交付 .pptx 文件时，加载 fcp-pptx-builder 技能，把这份 Markdown 转成可编辑的 PPT。

## 注意事项
- 一页只讲一个观点；放不下就拆页，不要缩小字号。
- 不要整段贴文字，讲的内容和页面上的字不要重复。
"""

SLIDES_TEMPLATE = """---
marp: true
paginate: true
---

# <核心结论：一句话>

<汇报人> · <日期>

---

## <这一页的结论>

- <要点 1>
- <要点 2>
- <要点 3>

> 数据来源：<…>

---

## 需要您决定

1. <决定事项 1>
2. <决定事项 2>
"""

SLIDES_SCRIPT = '''#!/usr/bin/env python3
"""把 Markdown 提纲转成页面骨架。用法：python3 outline_to_slides.py 提纲.md 输出.md

提纲里的一级标题作为封面，每个二级标题变成一页（标题即结论），下面的列表项作为这一页的要点。
"""

import pathlib
import sys


def convert(outline: str) -> str:
    pages, cover, current = [], "", None
    for line in outline.splitlines():
        text = line.strip()
        if text.startswith("# "):
            cover = text[2:].strip()
        elif text.startswith("## "):
            current = {"title": text[3:].strip(), "points": []}
            pages.append(current)
        elif text.startswith(("- ", "* ")) and current is not None:
            current["points"].append(text[2:].strip())
    slides = [f"# {cover or '<核心结论>'}"]
    for page in pages:
        body = "\\n".join(f"- {point}" for point in page["points"][:3]) or "- <要点>"
        slides.append(f"## {page['title']}\\n\\n{body}")
    header = "---\\nmarp: true\\npaginate: true\\n---\\n\\n"
    return header + "\\n\\n---\\n\\n".join(slides) + "\\n"


def main() -> int:
    if len(sys.argv) != 3:
        print("用法：python3 outline_to_slides.py 提纲.md 输出.md")
        return 2
    source, target = pathlib.Path(sys.argv[1]), pathlib.Path(sys.argv[2])
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(convert(source.read_text(encoding="utf-8")), encoding="utf-8")
    print(f"已生成 {target}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
'''

PPTX_SVG_SCRIPT = '''#!/usr/bin/env python3
"""把 SVG 转成 PNG，好插进 PPT（PowerPoint 和 python-pptx 都不能直接插 SVG）。

用法：python3 svg_to_png.py 输入.svg 输出.png [--scale 2]
依赖 resvg-py（写在技能的 requirements.txt 里，平台预装），不需要 cairo 等系统库。
"""
import argparse
import re
import sys
from pathlib import Path


def scaled(svg: str, scale: float) -> str:
    """按倍数放大：改根元素的 width/height，没有就从 viewBox 推出来。"""
    if scale == 1:
        return svg
    head = re.search(r"<svg\\b[^>]*>", svg)
    if not head:
        return svg
    tag = head.group(0)
    size = {}
    for name in ("width", "height"):
        found = re.search(rf\'\\b{name}="([0-9.]+)(px)?"\', tag)
        if found:
            size[name] = float(found.group(1))
    if len(size) < 2:
        box = re.search(r\'viewBox="[-0-9. ]*?([0-9.]+)[ ,]+([0-9.]+)"\', tag)
        if not box:
            return svg
        size = {"width": float(box.group(1)), "height": float(box.group(2))}
        if "viewBox" not in tag:
            return svg
    new = re.sub(r\'\\s(width|height)="[^"]*"\', "", tag)
    new = new.replace("<svg", f\'<svg width="{size["width"] * scale:g}" height="{size["height"] * scale:g}"\', 1)
    return svg.replace(tag, new, 1)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("source")
    parser.add_argument("target")
    parser.add_argument("--scale", type=float, default=1.0, help="放大倍数，投影用 2")
    args = parser.parse_args()
    try:
        import resvg_py
    except ImportError:
        print("没有 resvg_py：这个技能的依赖还没装好，请告诉用户（依赖写在技能的 requirements.txt）", file=sys.stderr)
        return 2
    svg = scaled(Path(args.source).read_text(encoding="utf-8"), args.scale)
    data = resvg_py.svg_to_bytes(svg_string=svg)
    Path(args.target).write_bytes(bytes(data))
    print(f"已生成 {args.target}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
'''

PPTX_SKILL = """---
name: fcp-pptx-builder
description: 在需要交付 PowerPoint 文件（.pptx）时使用：把 Markdown 页面稿一键转成可编辑的 PPT（文字、表格、讲稿备注都能再改），或用 python-pptx 直接生成、修改 PPT，交付前核对页数和内容
---

# 生成可编辑的 PPT

## 适用场景
- 用户要的是 .pptx 文件，而不只是提纲或 Markdown
- 已有一份 .pptx，要改其中的文字、表格或补几页

## 做法一：Markdown 页面稿转 PPT（首选）
1. 先写页面稿，比如 `slides/汇报.md`，写法见 `references/写法.md`。还没有提纲时，可以先用 fcp-markdown-slides 技能写提纲和页面。
2. 转换：`python3 <本技能目录>/scripts/md_to_pptx.py slides/汇报.md slides/汇报.pptx`
   - 公司有 PPT 模板时加 `--template 模板.pptx`，会沿用模板的母版和配色
3. 核对：`python3 <本技能目录>/scripts/check_pptx.py slides/汇报.pptx`
   - 会列出页数、每页标题，并提示空页、要点过多、文字过长的页
   - 有提示就回到页面稿改，再转一次
4. 交付时告诉用户：文件路径、页数、每页标题，以及需要他在 PowerPoint 或 WPS 里过一遍的地方。

## 做法二：直接用 python-pptx
要画原生图表、或者改一份已有的 PPT 时用。
- 改已有 PPT：先列出每页的版式、文本框、表格和图表，只改用户要求的内容，保留原有格式，另存为新文件，除非用户要求原地修改。
- 新建：版式用空白页，位置和大小用 `Inches` 或 `Cm`，字号用 `Pt`，都写明确的值。
- 脚本按用途命名，比如 `edit_pptx.py`。不要命名成 `pptx.py`、`inspect.py`，否则会遮住同名的库。

一个带图表的最小例子：

```python
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
title = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.8))
title.text_frame.text = "告警量连续三周下降"
title.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
data = CategoryChartData()
data.categories = ["第 1 周", "第 2 周", "第 3 周"]
data.add_series("告警数", [120, 86, 54])
slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.6), Inches(11.5), Inches(4.8), data)
prs.save("slides/告警趋势.pptx")
```

## 插 SVG 图（流程图、架构图、图标）
PowerPoint 和 python-pptx 都不能直接插 SVG，先转成 PNG 再插：

```bash
python3 <本技能目录>/scripts/svg_to_png.py 图/架构.svg 图/架构.png --scale 2
```

- 转换用的 resvg-py 写在本技能的 `requirements.txt` 里，平台按技能依赖预装，**不用自己安装**；SVG 里的中文用运行环境自带的中文字体
- `--scale 2` 出两倍分辨率，投影更清楚；插图用 `slide.shapes.add_picture("图/架构.png", left, top, width=...)`，只给宽度就会按比例缩放
- 提示没有 resvg_py：说明这个技能的依赖还没装好（刚加依赖、评审还没通过时会这样），如实告诉用户，不要改用别的办法自己装

## 交付前检查
- 重新打开文件，核对页数、每页标题、表格和图表的数据、讲稿备注
- 文字放不放得下：每页要点不超过 6 条，单条不超过 40 个字
- 原生图表里嵌的数据表也是交付内容，数字要和来源一致
- python-pptx 不是渲染引擎，上面的检查保证不了视觉效果，要提醒用户打开看一遍

## 注意事项
- 运行环境已经预装了 python-pptx，直接用，不要另行安装；本技能另外要的库写在 `requirements.txt`，平台会预装；真缺依赖时如实告诉用户
- 产出放在工作目录里（比如 `slides/`），不要写进技能目录
- 图片用本地文件，不要用网络图片地址

## 来源
改编自 DeepSeek Harness 的 office-pptx 技能（MIT 许可），许可全文见 `references/LICENSE-office-pptx.md`。
"""

PPTX_SYNTAX = """# 页面稿写法

一份页面稿就是一个 Markdown 文件，`md_to_pptx.py` 按下面的规则转成 PPT。

| 写法 | 变成 |
|---|---|
| `# 标题` | 封面页的大标题；紧跟的普通文字作为副标题 |
| `## 标题` | 新的一页，标题写这一页的结论 |
| `- 要点` | 项目符号；前面缩进 2 个空格是下一级，最多三级 |
| `1. 步骤` | 编号要点 |
| 普通文字 | 正文段落 |
| Markdown 表格 | PPT 原生表格，第一行是表头 |
| `> 文字` | 页面底部的小字，用来写数据来源 |
| `备注：文字` 或 `Note: 文字` | 演讲者备注，放映时观众看不到 |
| `**加粗**` | 加粗 |
| `---`、Marp 头部 | 忽略，可以照常保留 |

## 例子

```markdown
# 内部运维看板项目周报
第 43 周 · 平台部

## 告警量连续三周下降，误报减少是主因
- 本周告警 54 条，比上周少 37%
- 误报规则下线 6 条
  - 磁盘水位规则改为按增速告警
> 数据来源：运维看板告警统计
备注：这一页先讲结论，再讲原因

## 下周计划
| 事项 | 负责人 | 完成时间 |
|---|---|---|
| 接入日志告警 | 陈 | 周三 |
| 告警分级 | 罗 | 周五 |
```
"""

PPTX_LICENSE = """# office-pptx 许可

本技能改编自 DeepSeek Harness（https://github.com/deepseek-ai/deepseek-harness）中的
`packages/skill/skill-office/assets/office-pptx` 技能，原技能许可如下。

MIT License

Copyright (c) 2026 DeepSeek

Permission is hereby granted, free of charge, to any person obtaining a copy
of this software and associated documentation files (the "Software"), to deal
in the Software without restriction, including without limitation the rights
to use, copy, modify, merge, publish, distribute, sublicense, and/or sell
copies of the Software, and to permit persons to whom the Software is
furnished to do so, subject to the following conditions:

The above copyright notice and this permission notice shall be included in all
copies or substantial portions of the Software.

THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR
IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY,
FITNESS FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE
AUTHORS OR COPYRIGHT HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER
LIABILITY, WHETHER IN AN ACTION OF CONTRACT, TORT OR OTHERWISE, ARISING FROM,
OUT OF OR IN CONNECTION WITH THE SOFTWARE OR THE USE OR OTHER DEALINGS IN THE
SOFTWARE.
"""

PPTX_SCRIPT = r'''#!/usr/bin/env python3
"""把 Markdown 页面稿转成可编辑的 PowerPoint（.pptx）。

用法：python3 md_to_pptx.py 输入.md 输出.pptx [--template 模板.pptx]
写法见本技能的 references/写法.md。依赖 python-pptx（FCP 运行镜像已预装）。
"""

import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt
except ImportError:
    print("缺少 python-pptx：请让管理员在运行镜像里预装，或在技能依赖里声明 python-pptx", file=sys.stderr)
    sys.exit(3)

FONT = "Microsoft YaHei"
DARK = RGBColor(0x1F, 0x2D, 0x3D)
ACCENT = RGBColor(0x32, 0x71, 0xC8)
GREY = RGBColor(0x6B, 0x77, 0x85)
LIGHT = RGBColor(0xEE, 0xF3, 0xFA)
BULLET_SIZES = (20, 18, 16)
NOTE = re.compile(r"^(note|notes|备注)\s*[:：]\s*(.*)$", re.I)
BULLET = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
RULE = re.compile(r"^:?-{2,}:?$")
BOLD = re.compile(r"\*\*(.+?)\*\*")


def parse(text):
    lines = text.splitlines()
    if lines and lines[0].strip() == "---":  # Marp / YAML 头
        for index in range(1, len(lines)):
            if lines[index].strip() == "---":
                lines = lines[index + 1 :]
                break
    slides, current = [], None
    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("# "):
            current = {"cover": True, "title": stripped[2:].strip(), "items": [], "notes": [], "source": []}
            slides.append(current)
            continue
        if stripped.startswith("## "):
            current = {"cover": False, "title": stripped[3:].strip(), "items": [], "notes": [], "source": []}
            slides.append(current)
            continue
        if current is None or not stripped or stripped == "---" or stripped.startswith("<!--"):
            continue
        note = NOTE.match(stripped)
        if note:
            current["notes"].append(note.group(2))
        elif stripped.startswith(">"):
            current["source"].append(stripped.lstrip(">").strip())
        elif stripped.startswith("|"):
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            if all(RULE.match(cell) for cell in cells if cell):
                continue
            items = current["items"]
            if items and items[-1][0] == "table":
                items[-1][1].append(cells)
            else:
                items.append(("table", [cells]))
        else:
            bullet = BULLET.match(line)
            if bullet:
                level = min(len(bullet.group(1).replace("\t", "  ")) // 2, 2)
                numbered = bullet.group(2)[0].isdigit()
                label = (bullet.group(2) + " ") if numbered else ""
                current["items"].append(("bullet", (level, label + bullet.group(3))))
            else:
                current["items"].append(("text", stripped))
    return slides


def style(run, size, *, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    properties = run._r.get_or_add_rPr()
    east_asian = properties.find(qn("a:ea"))
    if east_asian is None:
        east_asian = properties.makeelement(qn("a:ea"), {})
        properties.append(east_asian)
    east_asian.set("typeface", FONT)


def write(paragraph, text, size, *, bold=False, color=DARK):
    """支持 **加粗** 的一段文字。"""
    position = 0
    for match in BOLD.finditer(text):
        if match.start() > position:
            style(paragraph.add_run(), size, bold=bold, color=color)
            paragraph.runs[-1].text = text[position : match.start()]
        style(paragraph.add_run(), size, bold=True, color=color)
        paragraph.runs[-1].text = match.group(1)
        position = match.end()
    if position < len(text) or not paragraph.runs:
        style(paragraph.add_run(), size, bold=bold, color=color)
        paragraph.runs[-1].text = text[position:]


def textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.text_frame.word_wrap = True
    return box.text_frame


def bar(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def cover(slide, spec, width, height):
    bar(slide, 0, 0, width, height, LIGHT)
    bar(slide, Inches(0.8), int(height * 0.36), Inches(0.12), Inches(1.3), ACCENT)
    frame = textbox(slide, Inches(1.1), int(height * 0.34), width - Inches(2), Inches(1.2))
    write(frame.paragraphs[0], spec["title"], 40, bold=True)
    subtitle = [value for kind, value in spec["items"] if kind == "text"]
    if subtitle:
        frame = textbox(slide, Inches(1.1), int(height * 0.34) + Inches(1.3), width - Inches(2), Inches(1))
        write(frame.paragraphs[0], " · ".join(subtitle), 20, color=GREY)


def content(slide, spec, width, height, number):
    frame = textbox(slide, Inches(0.6), Inches(0.35), width - Inches(1.2), Inches(0.9))
    write(frame.paragraphs[0], spec["title"], 28, bold=True)
    bar(slide, Inches(0.6), Inches(1.25), Inches(1.2), Inches(0.06), ACCENT)
    top, bottom = Inches(1.55), height - Inches(0.9)
    groups = []
    for kind, value in spec["items"]:
        if kind == "table":
            groups.append(("table", value))
        elif groups and groups[-1][0] == "text":
            groups[-1][1].append((kind, value))
        else:
            groups.append(("text", [(kind, value)]))
    for kind, value in groups:
        if top >= bottom:
            print(f"提示：第 {number} 页内容放不下，请拆页", file=sys.stderr)
            break
        if kind == "table":
            columns = max(len(row) for row in value)
            row_height = Inches(0.45)
            shape = slide.shapes.add_table(
                len(value), columns, Inches(0.6), top, width - Inches(1.2), row_height * len(value)
            )
            for r, row in enumerate(value):
                for c in range(columns):
                    cell = shape.table.cell(r, c)
                    cell.text_frame.paragraphs[0].text = ""
                    write(cell.text_frame.paragraphs[0], row[c] if c < len(row) else "", 14, bold=(r == 0),
                          color=RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else DARK)
            top += row_height * len(value) + Inches(0.25)
        else:
            lines = len(value)
            box_height = min(bottom - top, Inches(0.5) * lines + Inches(0.2))
            frame = textbox(slide, Inches(0.6), top, width - Inches(1.2), box_height)
            for index, (item_kind, item) in enumerate(value):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.space_after = Pt(8)
                if item_kind == "bullet":
                    level, text = item
                    paragraph.level = level
                    marker = "" if text[:1].isdigit() else ("• " if level == 0 else "– ")
                    write(paragraph, "    " * level + marker + text, BULLET_SIZES[level])
                else:
                    write(paragraph, item, 18)
            top += box_height + Inches(0.15)
    if spec["source"]:
        frame = textbox(slide, Inches(0.6), height - Inches(0.75), width - Inches(2), Inches(0.4))
        write(frame.paragraphs[0], "；".join(spec["source"]), 11, color=GREY)
    frame = textbox(slide, width - Inches(1.2), height - Inches(0.75), Inches(0.8), Inches(0.4))
    write(frame.paragraphs[0], str(number), 11, color=GREY)


def build(slides, output, template=None):
    prs = Presentation(template) if template else Presentation()
    if not template:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layouts = prs.slide_layouts
    blank = layouts[6] if len(layouts) > 6 else layouts[len(layouts) - 1]
    for number, spec in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank)
        if spec["cover"]:
            cover(slide, spec, prs.slide_width, prs.slide_height)
        else:
            content(slide, spec, prs.slide_width, prs.slide_height, number)
        if spec["notes"]:
            slide.notes_slide.notes_text_frame.text = "\n".join(spec["notes"])
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))


def main(argv):
    args = list(argv)
    template = None
    if "--template" in args:
        index = args.index("--template")
        template = args[index + 1] if index + 1 < len(args) else None
        del args[index : index + 2]
    if len(args) != 2 or (template is not None and not Path(template).is_file()):
        print("用法：python3 md_to_pptx.py 输入.md 输出.pptx [--template 模板.pptx]")
        return 2
    source, output = Path(args[0]), Path(args[1])
    slides = parse(source.read_text(encoding="utf-8"))
    if not slides:
        print("页面稿里没有找到 # 或 ## 标题，没有生成 PPT", file=sys.stderr)
        return 1
    build(slides, output, template)
    print(f"已生成 {output}，共 {len(slides)} 页")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
'''

PPTX_CHECK_SCRIPT = r'''#!/usr/bin/env python3
"""交付前核对 PPT：页数、每页标题，提示空页、要点过多和文字过长的页。

用法：python3 check_pptx.py 文件.pptx [--max-slides N]
只做内容核对，保证不了视觉效果；有超出 --max-slides 的页数时退出码为 1。
"""

import sys

try:
    from pptx import Presentation
except ImportError:
    print("缺少 python-pptx：请让管理员在运行镜像里预装", file=sys.stderr)
    sys.exit(3)


def main(argv):
    args = list(argv)
    limit = None
    if "--max-slides" in args:
        index = args.index("--max-slides")
        limit = int(args[index + 1])
        del args[index : index + 2]
    if len(args) != 1:
        print("用法：python3 check_pptx.py 文件.pptx [--max-slides N]")
        return 2
    prs = Presentation(args[0])
    warnings = 0
    print(f"共 {len(prs.slides)} 页")
    for number, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.extend(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                texts.append("[表格]")
        title = texts[0] if texts else "（空页）"
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            notes = "  [有备注]"
        print(f"{number:>2}. {title}{notes}")
        body = [t for t in texts[1:] if not t.isdigit()]
        if not texts:
            print("    提示：空页")
            warnings += 1
        if len(body) > 7:
            print(f"    提示：要点 {len(body)} 条，建议不超过 6 条，考虑拆页")
            warnings += 1
        for text in body:
            if len(text) > 60:
                print(f"    提示：这一条有 {len(text)} 个字，可能放不下：{text[:20]}…")
                warnings += 1
    print(f"提示 {warnings} 处" if warnings else "没有发现问题；视觉效果请在 PowerPoint 或 WPS 里过一遍")
    if limit is not None and len(prs.slides) > limit:
        print(f"页数超过上限 {limit}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
'''

ENGLISH_SKILL = """---
name: fcp-english-sentence-coach
description: 在想练英语表达、把中文句子翻成地道英文、或请人批改英文句子时使用：先让用户自己造句，再指出错因并给出地道说法和同类练习
---

# 英语句型教练

## 适用场景
用户想练英语写作或口语表达，或者给出一句英文请你批改。

## 步骤
1. 用户给的是中文：先不要直接翻译，给一个提示句型，请用户自己写英文。
2. 用户写了英文：逐句批改，每处问题按「原句 → 改后 → 错因」列出，错因用 `references/常见错误.md` 里的类别。
3. 给出 1～2 种更地道的说法，并说明适用场合（正式 / 口语）。
4. 针对出错的句型，再出 3 道同类的中译英小练习，等用户作答后再批改。
5. 最后用一句话总结这次最需要注意的一点。

## 注意事项
- 用中文解释，例句用英文；一次只纠正最重要的 3 处，不要把用户压垮。
- 不要只给答案：先让用户想，再讲为什么。
"""

ENGLISH_REFERENCE = """# 常见错误类别

| 类别 | 例子 | 怎么改 |
|---|---|---|
| 中式直译 | open the light | turn on the light |
| 时态 | I go there yesterday. | I went there yesterday. |
| 主谓一致 | He have two cats. | He has two cats. |
| 冠词 | I am engineer. | I am an engineer. |
| 介词 | discuss about the plan | discuss the plan |
| 可数不可数 | many informations | a lot of information |
| 句子不完整 | Because it rains. | We stayed home because it rained. |
| 用词正式度 | wanna / gonna 用在邮件里 | want to / going to |
"""


FILES_DIR = Path(__file__).with_name("showcase_files")


def _folder(name: str) -> dict[str, str]:
    """示范技能文件多（引入的现成技能、数据表）时放在 showcase_files/<名字>/ 下，按相对路径读进来。"""
    root = FILES_DIR / name
    return {
        path.relative_to(root).as_posix(): path.read_text(encoding="utf-8")
        for path in sorted(root.rglob("*"))
        if path.is_file() and "__pycache__" not in path.parts
    }


def _skill(
    name: str, title: str, summary: str, suitable: str, not_suitable: str, files: dict[str, str]
) -> dict[str, Any]:
    return {
        "kind": "Skill",
        "name": name,
        "title": title,
        "summary": summary,
        "suitable": suitable,
        "not_suitable": not_suitable,
        "files": files,
    }


def _sop(
    name: str, title: str, summary: str, suitable: str, not_suitable: str, body: str, try_it: str
) -> dict[str, Any]:
    return {
        "kind": "Sop",
        "name": name,
        "title": title,
        "summary": summary,
        "suitable": suitable,
        "not_suitable": not_suitable,
        "files": {"SOP.md": body},
        "layer": "L1",
        "try_it": try_it,
    }


SHOWCASE: list[dict[str, Any]] = [
    _sop(
        "fcp-tech-doc-sop",
        "技术文档编写流程（示范）",
        "定读者和范围 → 写初稿 → 评审定稿，三步写出别人能独立看懂的技术文档",
        "设计文档、方案说明、操作手册、接口说明",
        "一两句话就能说清的临时说明",
        DOC_SOP,
        "新建任务「写一份内部运维看板的部署说明」，执行规范里选这条流程，交给 Agent 执行；"
        "看 Agent 是否按第 1 步只出大纲、回报检查点，并加载了 fcp-tech-doc-writer 技能",
    ),
    _sop(
        "fcp-code-change-sop",
        "代码改动交付流程（示范）",
        "说清改什么 → 实现并自测 → 按清单自查、走评审，代码改动的最小闭环",
        "功能开发、缺陷修复、重构",
        "只改文档、不涉及代码的任务",
        CODE_SOP,
        "新建任务「给运维看板加一个健康检查接口」，绑这条流程，交给 Agent 执行；"
        "第 2 步可以看 Agent 怎么用 fde_task_complete 跑测试、按清单自查",
    ),
    _sop(
        "fcp-presentation-sop",
        "汇报材料（PPT）制作流程（示范）",
        "定受众和核心结论 → 写页面 → 生成 PPT 并走查定稿，交付一份结论先行、可编辑的 .pptx",
        "项目周报、方案汇报、复盘分享",
        "需要精美视觉设计的对外宣传材料",
        SLIDES_SOP,
        "新建任务「做一份运维看板的项目周报」，绑这条流程；看 Agent 先出提纲、再写页面，最后用 fcp-pptx-builder 生成 slides/汇报.pptx",
    ),
    _skill(
        "fcp-tech-doc-writer",
        "技术文档写作（示范）",
        "按读者和范围定结构、结论先行，附文档模板和自查清单",
        "设计文档、方案、手册等交给别人看的技术文档",
        "代码注释、提交说明这类很短的文字",
        {"SKILL.md": DOC_SKILL, "references/文档模板.md": DOC_TEMPLATE},
    ),
    _skill(
        "fcp-code-review-checklist",
        "代码评审清单（示范）",
        "提交前自查或评审改动：按清单逐项检查，并用脚本找出遗留的调试代码和待办",
        "合并请求之前的自查、评审别人的代码",
        "架构层面的设计评审",
        {"SKILL.md": REVIEW_SKILL, "scripts/find_leftovers.py": REVIEW_SCRIPT},
    ),
    _skill(
        "fcp-markdown-slides",
        "用 Markdown 写汇报 PPT（示范）",
        "先写提纲再写页面，每页标题即结论；脚本把提纲转成页面骨架，可转成 PPT",
        "项目汇报、方案讲解、复盘分享",
        "需要复杂动画和精美视觉的材料",
        {
            "SKILL.md": SLIDES_SKILL,
            "templates/页面模板.md": SLIDES_TEMPLATE,
            "scripts/outline_to_slides.py": SLIDES_SCRIPT,
        },
    ),
    _skill(
        "fcp-pptx-builder",
        "生成可编辑的 PPT（示范）",
        "把 Markdown 页面稿一键转成可编辑的 .pptx，或用 python-pptx 直接生成、修改 PPT，交付前核对页数和内容",
        "需要交付 .pptx 文件的汇报、周报、方案讲解；修改已有 PPT 的文字和表格",
        "需要复杂动画、精美视觉设计的对外宣传材料",
        {
            "SKILL.md": PPTX_SKILL,
            "references/写法.md": PPTX_SYNTAX,
            "references/LICENSE-office-pptx.md": PPTX_LICENSE,
            "scripts/md_to_pptx.py": PPTX_SCRIPT,
            "scripts/check_pptx.py": PPTX_CHECK_SCRIPT,
            "scripts/svg_to_png.py": PPTX_SVG_SCRIPT,
            # 技能依赖第 2 层：平台按这个文件在 Agent 机器上预装，resvg-py 自带渲染器，不需要系统库
            "requirements.txt": "resvg-py\n",
        },
    ),
    _skill(
        "fcp-english-sentence-coach",
        "英语句型教练（示范）",
        "先让你自己造句，再指出错因、给地道说法和同类练习",
        "练英语写作和表达、批改英文句子",
        "整篇文档的专业翻译",
        {"SKILL.md": ENGLISH_SKILL, "references/常见错误.md": ENGLISH_REFERENCE},
    ),
    # 人文类（2026-10-26）：现成技能和开源数据拼装，来源与许可写在各自的 LICENSE 文件里
    _skill(
        "fcp-solar-term-poems",
        "节气诗笺（示范）",
        "给个日期，说出节气和节日、配上古诗原文，再按格律写一首应景的新诗，可排成诗笺网页",
        "周会开场、节日问候、项目里程碑、团队氛围",
        "需要严格学术考证的诗词研究",
        _folder("fcp-solar-term-poems"),
    ),
    _skill(
        "fcp-sudongpo-perspective",
        "苏东坡视角（示范）",
        "用苏轼的旷达心态和生活美学陪你聊挫折、迷茫和生活趣味，回答里自然引用东坡诗词",
        "工作之余的人生困惑、想换个心态看问题",
        "技术问题、严谨的学术论证、需要专业帮助的心理问题",
        _folder("fcp-sudongpo-perspective"),
    ),
    _skill(
        "fcp-chengyu-jinnang",
        "成语锦囊（示范）",
        "用一个成语点出眼下的状况和该怎么做（釜底抽薪、抛砖引玉……），需要时再讲典故",
        "排查问题、做计划、会后复盘时给团队一个好记的说法",
        "正式文档里的措辞润色",
        _folder("fcp-chengyu-jinnang"),
    ),
]

TRY_SKILL = {
    "fcp-tech-doc-writer": "在私聊里说「帮我写一份内部运维看板的部署说明大纲」，看 Agent 是否加载这个技能、按模板出大纲",
    "fcp-code-review-checklist": "在私聊里说「帮我评审一下 xxx 这段代码」，看 Agent 是否按清单逐项检查",
    "fcp-markdown-slides": "在私聊里说「帮我做一份运维看板项目周报的 PPT 提纲」，看 Agent 是否先问受众、再按每页一个结论写",
    "fcp-pptx-builder": "在私聊里说「把运维看板项目周报做成 PPT 文件」，看 Agent 是否先写页面稿、再转成 .pptx 并核对页数；"
    "生成的文件在工作目录的 slides/ 下",
    "fcp-solar-term-poems": "在私聊里说「今天是什么节气？给团队写首诗做周会开场，排成诗笺」，"
    "看 Agent 是否先跑脚本拿节气和古诗原文、照抄原文，再写一首标「新作」的七绝，诗笺在工作目录的 poems/ 下",
    "fcp-sudongpo-perspective": "在私聊里说「用苏东坡的视角聊聊：项目延期被批评了，心里有点堵」，"
    "看回答是否先共情、自嘲破冰，再给具体可做的建议，并自然引用东坡诗词",
    "fcp-chengyu-jinnang": "在私聊里说「给我个锦囊：线上偶发报错，我已经加了三次重试还是不行」，"
    "看 Agent 是否只点一个成语（如釜底抽薪）就停下；再问「什么意思」，看是否四句话讲完典故",
    "fcp-english-sentence-coach": "在私聊里说「我想练英语：把『这个问题我们下周再讨论』翻成英文」，看 Agent 是否先让你自己写",
}
for _item in SHOWCASE:
    if _item["kind"] == "Skill":
        _item["try_it"] = TRY_SKILL[_item["name"]]
